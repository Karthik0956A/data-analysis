"""
Week 2 – Generate EDA Framework Report (DOCX)
Run AFTER generate_diagrams.py and eda_framework.py
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

os.makedirs("report", exist_ok=True)

# ── Helpers (same as Week 1) ──────────────────────────────────────────────────

def set_heading(doc, text, level=1, color=(108, 99, 255)):
    h = doc.add_heading(text, level=level)
    run = h.runs[0] if h.runs else h.add_run(text)
    run.font.color.rgb = RGBColor(*color)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    return h

def add_para(doc, text, bold=False, italic=False, size=11, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    return p

def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(86, 227, 159)
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.space_after = Pt(2)
    return p

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    for j, h in enumerate(headers):
        cell = hdr.cells[j]
        cell.text = h
        run = cell.paragraphs[0].runs[0]
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "6C63FF")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, row in enumerate(rows):
        r = table.rows[i + 1]
        for j, val in enumerate(row):
            r.cells[j].text = val
    return table

def insert_image(doc, path, caption, width=5.5):
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(width))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph(f"Figure: {caption}")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(9)
        cap.runs[0].font.color.rgb = RGBColor(120, 120, 120)
    else:
        doc.add_paragraph(f"[Image not found: {path}]")

# ── Build Document ────────────────────────────────────────────────────────────

def build_report():
    doc = Document()
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # Title
    title = doc.add_heading("Week 2 – EDA & Visualization Framework", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.color.rgb = RGBColor(67, 203, 255)

    sub = doc.add_paragraph("Customer Churn Prediction | Exploratory Data Analysis Plan")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(13)
    sub.runs[0].font.color.rgb = RGBColor(100, 100, 120)
    doc.add_page_break()

    # ── 1. Introduction to EDA ────────────────────────────────────────────────
    set_heading(doc, "1. Introduction to EDA", level=1)
    add_para(doc,
        "Exploratory Data Analysis (EDA) is the critical first analytical phase of any data science "
        "project. Coined by statistician John Tukey in 1977, EDA involves summarising data's main "
        "characteristics, often using visual methods, before any formal modelling begins. It acts as "
        "a dialogue between the analyst and the data — revealing structure, spotting anomalies, "
        "testing assumptions, and generating hypotheses that guide the modelling strategy.")
    add_para(doc,
        "In the context of the Customer Churn Prediction project, EDA serves three primary functions: "
        "(1) understanding the distribution and relationships of 80+ raw features, (2) identifying "
        "data quality issues that must be resolved before modelling, and (3) generating preliminary "
        "hypotheses about which customer behaviours most strongly predict churn.")

    set_heading(doc, "1.1 Why EDA Matters", level=2, color=(255, 101, 132))
    importance = [
        "Prevents garbage-in-garbage-out: catches data quality issues before they corrupt models.",
        "Guides feature selection: identifies highly correlated or low-variance features to drop.",
        "Informs model choice: skewed data → tree models; linear relationships → regression.",
        "Builds stakeholder trust: visualisations make data tangible and understandable.",
        "Surfaces business insights even before ML modelling is complete.",
    ]
    for item in importance:
        add_bullet(doc, item)

    insert_image(doc, "diagrams/eda_workflow.png",
                 "Step-by-Step EDA Workflow for Customer Churn Prediction")
    doc.add_paragraph()

    # ── 2. Data Types & Potential Features ───────────────────────────────────
    set_heading(doc, "2. Data Types & Potential Features", level=1)
    add_para(doc,
        "The Customer Churn dataset is expected to comprise approximately 80 raw variables across "
        "five data domains. Understanding the data type of each variable determines which exploration "
        "technique and visualisation method is appropriate.")

    headers = ["Data Domain", "Example Features", "Type", "Analysis Approach"]
    rows = [
        ["Demographics",    "Age, Gender, SeniorCitizen, Dependents",          "Categorical / Discrete",  "Bar charts, group comparisons"],
        ["Contract Info",   "ContractType, PaperlessBilling, PaymentMethod",   "Categorical (Nominal)",   "Frequency tables, churn rates"],
        ["Service Usage",   "MonthlyCharges, TotalCharges, DataUsageGB",       "Continuous (Float)",      "Histograms, KDE, box plots"],
        ["Engagement",      "Tenure, NumSupportCalls, NumProductsUsed",        "Discrete (Int)",          "Bar charts, violin plots"],
        ["Network Quality", "DropCallRate, AvgNetworkLatency, Outages",        "Continuous (Float)",      "Scatter, correlation heatmap"],
    ]
    add_table(doc, headers, rows)
    doc.add_paragraph()

    insert_image(doc, "diagrams/data_types_map.png",
                 "Data Types Map & Analysis Decision Framework")
    doc.add_paragraph()

    # ── 3. Exploration Techniques ─────────────────────────────────────────────
    set_heading(doc, "3. Data Exploration Techniques", level=1)

    set_heading(doc, "3.1 Univariate Analysis", level=2, color=(67, 203, 255))
    add_para(doc,
        "Univariate analysis examines each variable individually. The goal is to understand its "
        "distribution, central tendency, spread, and presence of outliers.")
    uni_techniques = [
        ("Histograms",       "Reveal shape of distribution (normal, skewed, bimodal).",
         "plt.hist(df['monthly_charges'], bins=30)"),
        ("KDE Plots",        "Smooth probability density estimate; better than histograms for continuous data.",
         "df['tenure'].plot.kde()"),
        ("Box Plots",        "Visualise median, IQR, and outliers simultaneously.",
         "sns.boxplot(y=df['total_charges'])"),
        ("Bar Charts",       "Frequency of categorical variables.",
         "df['contract'].value_counts().plot.bar()"),
        ("Value Counts",     "Exact category frequencies.",
         "df['internet_service'].value_counts(normalize=True)"),
    ]
    for name, desc, code in uni_techniques:
        set_heading(doc, f"• {name}", level=3, color=(86, 227, 159))
        add_para(doc, desc)
        add_code(doc, code)

    set_heading(doc, "3.2 Bivariate Analysis", level=2, color=(67, 203, 255))
    add_para(doc,
        "Bivariate analysis examines relationships between two variables, particularly between "
        "each feature and the target variable (churn).")
    biv_techniques = [
        ("Churn Rate by Category",  "Group by categorical feature → compute churn rate → bar chart.",
         "df.groupby('contract')['churn'].mean().plot.bar()"),
        ("Scatter Plots",           "Show relationship between two continuous features; colour by churn.",
         "sns.scatterplot(data=df, x='tenure', y='monthly_charges', hue='churn')"),
        ("Violin Plots",            "Distribution shape + summary statistics split by churn label.",
         "sns.violinplot(data=df, x='churn', y='monthly_charges')"),
        ("Point-Biserial Corr.",    "Correlation between binary target and continuous features.",
         "df.corr()['churn'].sort_values()"),
    ]
    for name, desc, code in biv_techniques:
        set_heading(doc, f"• {name}", level=3, color=(86, 227, 159))
        add_para(doc, desc)
        add_code(doc, code)

    set_heading(doc, "3.3 Multivariate Analysis", level=2, color=(67, 203, 255))
    add_para(doc,
        "Multivariate analysis considers interactions among three or more variables, revealing "
        "complex patterns not visible in simpler analyses.")
    multi_techniques = [
        ("Correlation Heatmap",   "Full correlation matrix of all numerical features.",
         "sns.heatmap(df.corr(), annot=True, cmap='coolwarm')"),
        ("Pair Plot",             "All-vs-all scatter matrix; coloured by churn.",
         "sns.pairplot(df[num_cols], hue='churn')"),
        ("PCA Biplot",            "Reduce to 2D; visualise cluster separation.",
         "from sklearn.decomposition import PCA; pca = PCA(2); X_pca = pca.fit_transform(X)"),
        ("Parallel Coordinates",  "Visualise all features simultaneously across observations.",
         "pd.plotting.parallel_coordinates(df_sample, 'churn')"),
    ]
    for name, desc, code in multi_techniques:
        set_heading(doc, f"• {name}", level=3, color=(86, 227, 159))
        add_para(doc, desc)
        add_code(doc, code)

    # ── 4. Handling Missing Data & Outliers ───────────────────────────────────
    set_heading(doc, "4. Handling Missing Data & Outliers", level=1)
    set_heading(doc, "4.1 Missing Value Strategy", level=2, color=(255, 101, 132))
    mv_rows = [
        ["MCAR (Missing Completely At Random)", "< 5%",  "Simple median/mode imputation"],
        ["MAR (Missing At Random)",             "5–30%", "MICE (Multiple Imputation by Chained Equations)"],
        ["MNAR (Missing Not At Random)",        "Any",   "Flag indicator column + domain imputation"],
        ["Structural Missing",                  "Any",   "Treat as a separate category (e.g., 'No Service')"],
    ]
    add_table(doc, ["Missing Pattern", "Threshold", "Strategy"], mv_rows)
    doc.add_paragraph()

    set_heading(doc, "4.2 Outlier Detection & Treatment", level=2, color=(255, 101, 132))
    add_para(doc, "We use multiple methods to detect outliers, applying different treatments based on context:")
    out_rows = [
        ["IQR Method",         "Values outside [Q1-1.5×IQR, Q3+1.5×IQR]", "Cap to fence (Winsorisation)"],
        ["Z-Score",            "|Z| > 3",                                    "Remove or cap"],
        ["Isolation Forest",   "Unsupervised anomaly detection",             "Flag as anomaly; investigate"],
        ["Domain Knowledge",   "Business rules (e.g., tenure > 120 months)", "Remove or correct"],
    ]
    add_table(doc, ["Method", "Detection Criterion", "Treatment"], out_rows)
    doc.add_paragraph()

    # ── 5. Visualisation Strategies ───────────────────────────────────────────
    set_heading(doc, "5. Visualisation Strategies", level=1)
    add_para(doc,
        "A structured visualisation plan ensures every analytical question has a corresponding "
        "chart. The strategy matrix below maps analysis goals to specific plot types and Python "
        "implementation tools.")

    insert_image(doc, "diagrams/visualization_types.png",
                 "Visualization Strategy Matrix – Four Categories of EDA Plots")
    doc.add_paragraph()
    insert_image(doc, "diagrams/eda_sample_plots.png",
                 "Sample EDA Output Plots Generated on Synthetic Churn Dataset", width=6.0)
    doc.add_paragraph()

    # ── 6. Tools & Libraries ──────────────────────────────────────────────────
    set_heading(doc, "6. Tools & Python Libraries", level=1)
    lib_rows = [
        ["pandas",      "2.2.x", "Data loading, cleaning, groupby, pivot tables, missing value handling"],
        ["numpy",       "1.26.x","Numerical computation, array operations, statistical functions"],
        ["matplotlib",  "3.9.x", "Low-level plotting; full customisation; save high-DPI PNGs"],
        ["seaborn",     "0.13.x","Statistical plot library built on matplotlib; heatmaps, violin, pair"],
        ["plotly",      "5.22.x","Interactive dashboards; scatter, sunburst, Sankey, parallel coords"],
        ["scipy",       "1.13.x","Statistical tests (chi-squared, t-test, ANOVA) for EDA validation"],
        ["scikit-learn","1.5.x", "PCA, Isolation Forest, preprocessing (StandardScaler, encoder)"],
        ["missingno",   "0.5.x", "Specialised visualisation of missing value patterns"],
    ]
    add_table(doc, ["Library", "Version", "EDA Role"], lib_rows)
    doc.add_paragraph()

    # ── 7. Reporting & Documentation Plan ────────────────────────────────────
    set_heading(doc, "7. Reporting & Documentation Plan", level=1)
    add_para(doc,
        "All EDA findings will be documented in a structured Jupyter Notebook and summarised in "
        "a concise EDA Report. The report will follow this template:")
    sections = [
        "Executive Summary (1 page): Key findings, class balance, top 3 patterns discovered.",
        "Dataset Overview: Shape, data types, coverage dates, source systems.",
        "Data Quality Report: Missing value counts, outlier flags, duplicate rows.",
        "Univariate Analysis: Histograms/KDE for each numerical feature; bar charts for categorical.",
        "Bivariate Analysis: Churn rate by key features; scatter plots; violin plots.",
        "Multivariate Analysis: Correlation heatmap; PCA projection; pair plot subset.",
        "Key Hypotheses: Numbered list of hypotheses generated for modelling phase.",
        "Data Preparation Recommendations: Feature engineering ideas, encoding strategies.",
        "Appendix: Full correlation matrix; distribution plots for all 80 features.",
    ]
    for s in sections:
        add_bullet(doc, s)

    # ── 8. Conclusion ─────────────────────────────────────────────────────────
    set_heading(doc, "8. Conclusion", level=1)
    add_para(doc,
        "This Week 2 framework establishes a comprehensive, systematic approach to EDA that is "
        "adaptable to any structured tabular dataset. By combining univariate, bivariate, and "
        "multivariate techniques with a rich library of Python visualisation tools, the framework "
        "ensures no important pattern goes undetected before modelling begins. The structured "
        "documentation plan ensures findings are communicated clearly to both technical and "
        "non-technical stakeholders, setting the stage for informed, hypothesis-driven modelling "
        "in Week 3.")

    out_path = "report/Week2_EDA_Framework.docx"
    doc.save(out_path)
    print(f"✓ Report saved → {out_path}")


if __name__ == "__main__":
    build_report()
